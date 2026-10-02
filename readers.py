"""Читає вхідні файли (pdf, docx, txt, md, csv) і повертає текст для агентів."""
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def read_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    if ext == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".csv":
        # Для великих таблиць відправляємо лише початок, щоб не роздути запит
        lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
        head = "\n".join(lines[:60])
        note = f"\n... (усього рядків: {len(lines)})" if len(lines) > 60 else ""
        return head + note
    return path.read_text(encoding="utf-8", errors="replace")


def read_inputs(folder: Path) -> str:
    chunks = []
    for f in sorted(folder.iterdir()):
        if f.is_file() and f.suffix.lower() in {".pdf", ".docx", ".txt", ".md", ".csv"}:
            chunks.append(f"=== ФАЙЛ: {f.name} ===\n{read_file(f)}")
    if not chunks:
        raise SystemExit(f"У папці {folder} немає файлів. Покладіть туди методичку/дані.")
    return "\n\n".join(chunks)

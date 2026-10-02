"""Збирає фінальний Word-документ за налаштуваннями з config.py."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import config

FIGS_DIR = Path("output/figs")  # задається в build()


def _style_run(run, bold=False, size=None):
    run.font.name = config.FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), config.FONT)
    run.font.size = Pt(size or config.FONT_SIZE)
    run.bold = bold


def _para(doc, text="", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, mono=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = config.LINE_SPACING
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = Cm(config.FIRST_LINE_INDENT_CM)
    r = p.add_run(text)
    _style_run(r, bold=bold)
    if mono:
        r.font.name = "Consolas"
        r.font.size = Pt(11)
    return p


def _title_page(doc, plan, title_page):
    c = WD_ALIGN_PARAGRAPH.CENTER
    right = WD_ALIGN_PARAGRAPH.RIGHT
    t = title_page
    _para(doc, t["university"], align=c, indent=False)
    _para(doc, t["department"], align=c, indent=False)
    for _ in range(6):
        _para(doc, "", indent=False)
    _para(doc, f"ЗВІТ\nпро виконання лабораторної роботи № {plan.get('lab_number', '')}",
          bold=True, align=c, indent=False)
    _para(doc, f"з дисципліни «{plan.get('subject', '')}»", align=c, indent=False)
    _para(doc, f"Тема: {plan.get('title', '')}", align=c, indent=False)
    for _ in range(5):
        _para(doc, "", indent=False)
    _para(doc, t["student"], align=right, indent=False)
    _para(doc, t["teacher"], align=right, indent=False)
    for _ in range(4):
        _para(doc, "", indent=False)
    _para(doc, t["city_year"], align=c, indent=False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


FIG_RE = re.compile(r"\[РИС:\s*([^|\]]+?)\s*\|\s*([^\]]+?)\s*\]")
_fig_counter = 0


def _figure(doc, filename, caption):
    global _fig_counter
    path = FIGS_DIR / filename
    if not path.exists():
        _para(doc, f"[РИСУНОК {filename} не знайдено]", indent=False)
        return
    _fig_counter += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(14))
    _para(doc, f"Рисунок {_fig_counter} – {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def _body(doc, text):
    in_code = False
    for line in text.splitlines():
        fig = FIG_RE.search(line)
        if fig and not in_code:
            _figure(doc, fig.group(1), fig.group(2))
        elif line.strip().startswith("```"):
            in_code = not in_code
        elif in_code:
            _para(doc, line, indent=False, mono=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.strip().startswith("- "):
            _para(doc, "– " + line.strip()[2:])
        elif line.strip():
            _para(doc, line.strip())


def build(plan: dict, sections: dict, out_path, title_page: dict | None = None, figs_dir=None):
    global _fig_counter, FIGS_DIR
    if figs_dir:
        FIGS_DIR = Path(figs_dir)
    _fig_counter = 0
    doc = Document()
    s = doc.sections[0]
    s.left_margin, s.right_margin = Cm(config.MARGINS_CM["left"]), Cm(config.MARGINS_CM["right"])
    s.top_margin, s.bottom_margin = Cm(config.MARGINS_CM["top"]), Cm(config.MARGINS_CM["bottom"])
    _title_page(doc, plan, {**config.TITLE_PAGE, **(title_page or {})})
    for name, text in sections.items():
        _para(doc, name, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
        _body(doc, text)
    doc.save(out_path)
